# Звіт до роботи КП 2 (структура - за розд. 3 матеріалів до роботи): -> ТВ-43.КП_2.СтепаненкоНЮ.docx
# Протокол тестування формується зі справжніх запусків krok_parse.py.
import contextlib, io, os, sys, textwrap
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, 'tests'))
import krok_parse
from run_tests import CASES          # той самий набір випадків, що й в автоматичних тестах

OUT = os.path.join(HERE, 'ТВ-43.КП_2.СтепаненкоНЮ.docx')
RULES = open(os.path.join(HERE, 'krok_parse.py'), encoding='utf-8').read().count('\n@rule\n')
GRAMMAR = os.path.join(HERE, '..', '..', 'RGR', 'RGR1', 'grammar', 'krok.ebnf')
doc = Document()
J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER


def font(run, name='Times New Roman', size=12, bold=False, italic=False):
    run.font.name, run.font.size, run.bold, run.italic = name, Pt(size), bold, italic
    rp = run._element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    return run


def para(text, bold=False, size=12, align=J, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.15
    pf.first_line_indent = Cm(1.0) if indent else Cm(0)
    p.alignment = align
    for i, part in enumerate(text.split('`')):
        if part:
            font(p.add_run(part), 'Consolas' if i % 2 else 'Times New Roman', size - (1 if i % 2 else 0), bold)
    return p


def heading(text):
    p = para(text, bold=True, size=13, align=L, indent=False)
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(4)
    p.paragraph_format.keep_with_next = True


def code(lines, size=8):
    for line in [w for l in lines for w in (textwrap.wrap(l, 100, subsequent_indent='    ') if len(l) > 100 else [l])]:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing, pf.left_indent = Pt(0), Pt(0), 1.0, Cm(0.3)
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F2F2F2')
        p._p.get_or_add_pPr().append(sh)
        font(p.add_run(line.replace('\t', '    ') or ' '), 'Consolas', size)
    doc.add_paragraph().paragraph_format.line_spacing = 0.5


def table(rows, widths, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).width = Cm(widths[c])
            p = t.cell(r, c).paragraphs[0]
            p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(0), 1.0
            font(p.add_run(str(text)), size=size, bold=(r == 0))
    doc.add_paragraph().paragraph_format.line_spacing = 0.5


def run_parser(text):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        ok = krok_parse.parse(text)[1]
    out = buf.getvalue().rstrip('\n').split('\n')
    start = next(i for i, l in enumerate(out) if l.startswith('Lexer:'))
    return ok, out[start:]                       # лістинг лексем КП1 у протоколі не повторюємо


# ---------------------------------------------------------------- оформлення
sec = doc.sections[0]
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(1.5), Cm(2), Cm(2)
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run('Степаненко Назар, група ТВ-43, робота КП № 2'), size=10, italic=True)
fp = sec.footer.paragraphs[0]; fp.alignment = C
r = fp.add_run()
for kind in ('begin', 'instr', 'end'):
    el = OxmlElement('w:instrText' if kind == 'instr' else 'w:fldChar')
    if kind == 'instr':
        el.text = 'PAGE'
    else:
        el.set(qn('w:fldCharType'), kind)
    r._r.append(el)
font(r, size=10)

para('Синтаксичний аналіз методом рекурсивного спуску', bold=True, size=16, align=C, indent=False)
para('Мета роботи — програмна реалізація синтаксичного аналізатора (парсера) мови Krok методом '
     'рекурсивного спуску.')

heading('1. Коротка характеристика мови')
para('Krok — проста імперативна процедурна мова програмування загального призначення, статично і суворо '
     'типізована; мова-прототип — Go, рівень складності 2. Програма складається з одного файлу: `package main`, '
     '`import "fmt"`, оголошення змінних, констант, типів і функцій. Мова має типи `int`, `float`, `bool`, '
     '`string`, користувацькі скалярні типи, масиви та структури; інструкції присвоєння, розгалуження '
     '`if / else if / else`, циклу `for` (з передумовою та трикомпонентного), введення й виведення через '
     'пакет `fmt`. Крапку з комою наприкінці рядка вставляє лексичний аналізатор.')

heading('2. Повна граматика мови')
para('Граматика належить до класу LL(1) і не містить лівої рекурсії: ліва асоціативність бінарних операторів '
     'виражена повтореннями `{ ... }`, права асоціативність `**` — правою рекурсією через `UnaryExpr`. '
     'Тому кожне правило безпосередньо перетворюється на функцію рекурсивного спуску. Синтаксична частина '
     'граматики (аксіома `Program`):')
text = open(GRAMMAR, encoding='utf-8').read()
code(text[text.index('Program  '):text.index('(* ====', text.index('Program  '))].rstrip().split('\n'))
para('Лексична частина граматики (`Ident`, `IntLit`, `FloatLit`, `StringLit` та ін.) реалізована лексичним '
     'аналізатором КП 1; для парсера ці нетермінали є токенами `id`, `intnum`, `floatnum`, `strlit`, `boolval`.')

heading('3. Лексичний аналізатор')
para('Використано лексичний аналізатор `krok_lex.py` з роботи КП 1, побудований методом діаграми станів. '
     'Парсер викликає функцію `analyze()` і отримує таблицю розбору `tableOfSymb` у форматі '
     '`{ n_rec : (num_line, lexeme, token, idxIdConst) }`. Синтаксичний аналіз виконується лише тоді, коли '
     'лексичний аналіз завершився успішно. Крапки з комою, вставлені лексичним аналізатором наприкінці рядків, '
     'потрапляють у таблицю розбору як звичайні токени `(;, punct)`.')

heading('4. Програмна реалізація синтаксичного аналізатора')
for line in (
    'Парсер `krok_parse.py` побудовано за зразком `parse_02.py`. Глобальна змінна `numRow` — номер поточного '
    'запису таблиці розбору; `getSymb()` повертає номер рядка, лексему й токен поточного запису; '
    '`parseToken(lexeme, token)` перевіряє, що поточний запис збігається з очікуваним, і переходить до '
    'наступного; `failParse(str, tuple)` виводить діагностичне повідомлення і аварійно завершує роботу.',
    'Нетерміналам синтаксичної граматики відповідають функції `parse<Нетермінал>()` — усього {0}. Окремих '
    'функцій не мають лише правила, що складаються з одного токена (`RelOp`, `AddOp`, `MulOp`, `UnaryOp`, '
    '`AssignOp`, `IncDecOp`, `BoolLit`, `TypeName = Ident`): їх перевіряє безпосередньо порівняння поточного '
    'токена. Альтернативу обирає поточна лексема (граматика LL(1)), повторення `{{ ... }}` реалізовано циклом '
    '`while`, необов\'язкову частину `[ ... ]` — умовою `if`.'.format(RULES),
    'Відмінності від зразка: (1) функція `peek()` повертає поточну лексему й токен без аварійного завершення '
    'в кінці таблиці — вона потрібна для вибору альтернатив і повторень; (2) `parseToken(None, token)` приймає '
    'будь-яку лексему заданого токена (ідентифікатор, літерал), а `expect(lexeme)` бере токен з таблиці лексем '
    'мови; (3) відступи у трасі розбору й виведення імені функції виконує декоратор `@rule`, а не явні виклики '
    '`nextIndt()` і `predIndt()` у кожній функції.',
    'Коди аварійного завершення: 1 — невідповідність токенів; 2 — невідповідність інструкцій; 3 — помилка у '
    'виразі; 4 — помилка у типі; 5 — невідповідність оголошень верхнього рівня; 1001, 1002 — неочікуваний '
    'кінець програми.'):
    para(line)

heading('5. План тестування')
groups = {}
for g, name, src, exp in CASES:
    groups.setdefault(g, [0, exp])[0] += 1
expect = {True: 'успішний розбір', False: 'помилка, локалізація (номер рядка), діагностичне повідомлення'}
rows = [('№', 'Тип випробування', 'Очікуваний результат', 'Кількість')]
rows += [(g.split()[0], g.split(' ', 1)[1], expect[e], n) for g, (n, e) in groups.items()]
para('Таблиця 1 — План тестування', align=L, indent=False, size=11)
table(rows, [0.8, 5.0, 9.0, 1.8])
para('Усі випадки виконуються автоматично програмою `tests/run_tests.py`; результат кожного додатково '
     'звіряється з незалежним розпізнавачем, що працює безпосередньо за граматикою специфікації мови. '
     'Обидва аналізатори ухвалили однакові рішення в усіх {0} випадках.'.format(len(CASES)))

heading('6. Протокол тестування')
COMMENTS = {
    '1': 'відсутній термінал виявлено в тому місці, де він очікувався; повідомлення вказує номер рядка, '
         'знайдену й очікувану лексеми.',
    '2': 'зайва лексема не відповідає жодному правилу на цій позиції — розбір аварійно завершено.',
    '3': 'лексичний аналізатор розпізнав помилково написане ключове слово як ідентифікатор, а парсер виявив, '
         'що на цьому місці ідентифікатор неприпустимий.',
    '4': 'недопустиму послідовність операторів виявлено; `a < b < c` відхилено, бо оператори порівняння '
         'неасоціативні.',
    '5': 'у трасі розбору множення вкладене глибше за додавання (`parseMulExpr` усередині `parseAddExpr`), '
         'тобто виконується раніше — пріоритет операторів дотримано.',
    '6': 'для `a - b - c` обидва віднімання розібрано циклом одного виклику `parseAddExpr` (ліва '
         'асоціативність); для `2 ** 3 ** 2` другий `**` вкладено у праву частину першого (права асоціативність).',
}
shown = [c for c in CASES if c[0][0] in '123456'] + [c for c in CASES if c[1] == 'nested.krok']
for k, (group, name, src, exp) in enumerate(shown, 1):
    para('Тест № {0} ({1}): {2}.'.format(k, group, name), bold=True, align=L, indent=False)
    lines = src.rstrip('\n').split('\n')
    if name != 'nested.krok':
        code(['{0:>2}  {1}'.format(i, l) for i, l in enumerate(lines, 1)])
    ok, out = run_parser(src)
    if ok and name != 'nested.krok':           # для виразів показуємо лише розбір інструкції
        a = next(i for i, l in enumerate(out) if 'parseSimpleStmt' in l)
        b = next(i for i in range(a + 1, len(out)) if 'parseStatementTail' in out[i])
        out = ['...'] + out[a:b] + ['...', out[-1]]
    elif ok:
        out = out[:3] + ['... ({0} рядків пропущено) ...'.format(len(out) - 8)] + out[-5:]
    else:
        out = ['...'] + out[-7:]
    code(['>python krok_parse.py test.krok'] + out)
    para('Оцінка: ' + (COMMENTS.get(group[0]) or 'вкладені цикли, розгалуження й літерали масивів розібрано '
                                                  'успішно; глибина вкладеності видна з відступів у трасі.'))

heading('7. Висновки')
for line in (
    'Реалізовано синтаксичний аналізатор мови Krok методом рекурсивного спуску: {0} функцій розбору за '
    'правилами синтаксичної граматики. Аналізатор працює з таблицею розбору лексичного аналізатора '
    'КП 1, розбирає всі конструкції мови й локалізує синтаксичні помилки з точністю до рядка.'.format(RULES),
    'Завдяки тому, що граматику ще на етапі специфікації було приведено до класу LL(1) і звільнено від лівої '
    'рекурсії, жодних додаткових трансформацій не знадобилося: для вибору альтернативи завжди достатньо однієї '
    'поточної лексеми, а розбір відбувається без повернень.',
    'Найцікавішим виявився розбір простої інструкції: присвоєння, оголошення з ініціалізацією, інкремент і '
    'виклик функції починаються з виразу, тож парсер спершу розбирає вираз і лише потім, за наступною '
    'лексемою, з\'ясовує, що це за інструкція. Перевірку того, чи є ліва частина присвоєння l-значенням, '
    'відкладено до семантичного аналізу (КП 3).',
    'Коректність підтверджено {0} тестами за планом із табл. 1 та звіркою з незалежним розпізнавачем, що працює '
    'безпосередньо за граматикою специфікації: рішення обох аналізаторів збіглися в усіх випадках.'.format(len(CASES))):
    para(line)

doc.save(OUT)
print('written', os.path.basename(OUT))
