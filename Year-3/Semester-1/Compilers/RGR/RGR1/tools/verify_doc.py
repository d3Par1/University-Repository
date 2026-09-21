"""Post-build verification of the generated documents.

Checks formatting of every .docx in the project folder (fonts, sizes, line spacing,
indents, margins, page numbering) and the exported .pdf (leftover build directives,
unresolved fields, images, table of contents).

Usage:  python tools/verify_doc.py
"""
import glob
import re
import sys
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
BODY_FONT, MONO_FONT = "Times New Roman", "Consolas"
# ДСТУ 3008: ліве ≥ 20 мм, праве ≥ 10 мм, верхнє і нижнє ≥ 20 мм
MIN_MARGINS_CM = {"left_margin": 2.0, "right_margin": 1.0, "top_margin": 2.0, "bottom_margin": 2.0}
LEFTOVERS = ["[@", "{#", "@rules", "@code", "@tokentable", "@checks", "@fullgrammar",
             "@grammarstats", "@transforms", "Оновіть поле", "None", "Error!"]


def iter_runs(doc):
    for p in doc.paragraphs:
        yield p, p.runs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p, p.runs


def check_docx(path):
    issues, notes = [], []
    d = Document(path)

    sec = d.sections[0]
    for attr, low in MIN_MARGINS_CM.items():
        cm = getattr(sec, attr).cm
        if cm + 1e-6 < low:
            issues.append(f"поле {attr} = {cm:.2f} см < {low} см (ДСТУ 3008)")
    notes.append("поля, см: " + ", ".join(f"{a.split('_')[0]} {getattr(sec, a).cm:.1f}"
                                          for a in MIN_MARGINS_CM))
    notes.append(f"різна перша сторінка (без номера): {sec.different_first_page_header_footer}")

    # Word drops direct formatting equal to the style, so a run without an explicit
    # font inherits it from its paragraph style — that is not a defect.
    fonts, sizes = Counter(), Counter()
    for p, runs in iter_runs(d):
        for r in runs:
            if not r.text.strip():
                continue
            fonts[r.font.name or f"{p.style.name} (успадковано)"] += 1
            sizes[r.font.size.pt if r.font.size else None] += 1
    for name, n in fonts.items():
        if name not in (BODY_FONT, MONO_FONT) and "успадковано" not in name:
            issues.append(f"сторонній шрифт {name!r} у {n} фрагментах")
    for style in {s for s in fonts if "успадковано" in s}:
        st = d.styles[style.split(" (")[0]]
        name = st.font.name
        while name is None and st.base_style is not None:   # walk the style chain
            st = st.base_style
            name = st.font.name
        name = name or d.styles["Normal"].font.name
        if name not in (BODY_FONT, MONO_FONT):
            issues.append(f"стиль {style.split(' (')[0]}: успадкований шрифт {name!r}")
    notes.append("шрифти: " + ", ".join(f"{k} — {v}" for k, v in fonts.most_common()))
    notes.append("кеглі: " + ", ".join(f"{k} — {v}" for k, v in sizes.most_common(8)))

    normal = d.styles["Normal"].paragraph_format
    notes.append(f"стиль Normal: інтервал {normal.line_spacing}, абзацний відступ "
                 f"{normal.first_line_indent.cm if normal.first_line_indent else 0:.2f} см")
    if normal.line_spacing != 1.5:
        issues.append(f"міжрядковий інтервал стилю Normal = {normal.line_spacing}, очікується 1.5")
    if not normal.first_line_indent or abs(normal.first_line_indent.cm - 1.25) > 0.01:
        issues.append("абзацний відступ стилю Normal ≠ 1.25 см")

    body = [p for p in d.paragraphs if p.style.name == "Normal" and p.text.strip()
            and p.paragraph_format.line_spacing not in (1.0, 0.6)]
    odd = [p.text[:40] for p in body if p.paragraph_format.line_spacing not in (None, 1.5)]
    if odd:
        issues.append(f"абзаци з нетиповим інтервалом: {odd[:3]}")

    footer = sec.footer.paragraphs[0]
    if "PAGE" not in footer._p.xml:
        issues.append("у колонтитулі немає поля PAGE (нумерації сторінок)")

    heads = Counter(p.style.name for p in d.paragraphs if p.style.name.startswith("Heading"))
    notes.append("заголовки: " + ", ".join(f"{k} — {v}" for k, v in sorted(heads.items())))
    return issues, notes


def check_pdf(path):
    issues, notes = [], []
    doc = pymupdf.open(path)
    text = "".join(p.get_text() for p in doc)
    for bad in LEFTOVERS:
        n = text.count(bad)
        if n:
            issues.append(f"у тексті залишилось {bad!r} ({n} разів)")
    images = sum(len(p.get_images()) for p in doc)
    notes.append(f"сторінок: {len(doc)}, зображень (діаграм): {images}")

    empty = [i + 1 for i, p in enumerate(doc) if not p.get_text().strip() and not p.get_images()]
    if empty:
        issues.append(f"порожні сторінки: {empty}")

    toc_pages = [i for i in range(min(6, len(doc))) if "ЗМІСТ" in doc[i].get_text()]
    if not toc_pages:
        issues.append("не знайдено сторінку «ЗМІСТ»")
    else:
        entries = re.findall(r"\.{2,}\s*(\d+)", "".join(doc[i].get_text() for i in range(1, 6)))
        if not entries:
            issues.append("у змісті немає номерів сторінок (поле не оновлене)")
        else:
            notes.append(f"пунктів змісту з номерами: {len(entries)}")

    numbered = sum(1 for p in doc
                   if any(b[4].strip().isdigit() for b in p.get_text("blocks")
                          if b[1] > p.rect.height * 0.90))
    if numbered < len(doc) - 1:          # титульний аркуш без номера — за задумом
        issues.append(f"номери сторінок лише на {numbered} із {len(doc)}")
    notes.append(f"сторінок із номером унизу: {numbered} із {len(doc)} (титулка без номера)")

    allowed = ("TimesNewRoman", "Consolas", "CambriaMath", "Cambria")
    used = {f[3].split("+")[-1] for p in doc for f in p.get_fonts()}
    alien = [f for f in used if not any(a in f for a in allowed)]
    if alien:
        issues.append(f"сторонні шрифти у PDF: {alien}")
    notes.append("шрифти PDF: " + ", ".join(sorted(used)))
    if "Степаненко Назар Юрійович" not in doc[0].get_text():
        issues.append("на титульному аркуші немає ПІБ виконавця")
    return issues, notes


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ok = True
    for docx in sorted(glob.glob(str(ROOT / "*.docx"))):
        pdf = Path(docx).with_suffix(".pdf")
        print("=" * 70)
        print(Path(docx).name)
        for label, (issues, notes) in (("DOCX", check_docx(docx)),
                                       ("PDF ", check_pdf(pdf) if pdf.exists() else ([f"немає {pdf.name}"], []))):
            for n in notes:
                print(f"  [{label}] {n}")
            for i in issues:
                ok = False
                print(f"  [{label}] ПРОБЛЕМА: {i}")
    print("=" * 70)
    print("формат документів: усе гаразд" if ok else "формат документів: є зауваження")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
