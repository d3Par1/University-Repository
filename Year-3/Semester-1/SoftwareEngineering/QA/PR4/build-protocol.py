"""Builds Протокол_ПР4_Степаненко.docx — bug reports for module "Агент БПЛА" (individual work).

Bugs come from bugs.json (also used for the Jira issues and the Excel report);
story texts come from ../PR1/backlog.json.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
BUGS = json.loads((HERE / "bugs.json").read_text(encoding="utf-8"))
BACKLOG = json.loads((HERE.parent / "PR1" / "backlog.json").read_text(encoding="utf-8"))
SHOTS = HERE / "screenshots"
FONT = "Times New Roman"
STORY_KEY_OFFSET = 13
C, R, J, L = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT,
              WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT)

doc = Document()
sec = doc.sections[0]
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
sec.different_first_page_header_footer = True


def set_font(style, size, bold=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)


set_font(doc.styles["Normal"], 14)
doc.styles["Normal"].paragraph_format.space_after = Pt(0)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, before, after in (("Heading 1", 15, 14, 6), ("Heading 2", 14, 10, 4)):
    st = doc.styles[name]
    set_font(st, size, bold=True, color=RGBColor(0, 0, 0))
    st.font.italic = False
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.15
set_font(doc.styles["List Bullet"], 14)


def field(paragraph, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


footer = sec.footer.paragraphs[0]
footer.alignment = C
field(footer, "PAGE")


def para(text="", align=J, size=None, bold=False, after=0, indent=True, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if indent and align == J:
        p.paragraph_format.first_line_indent = Cm(1.25)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    return p


def h1(text):
    return doc.add_heading(text, level=1)


def h2(text):
    return doc.add_heading(text, level=2)


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = J
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def row_flag(row, tag):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fill(cell, text, size, bold=False, align=L):
    cell.text = ""
    lines = text if isinstance(text, list) else [text]
    for n, line in enumerate(lines):
        p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(line).replace("UAVS-", "UAVS‑").replace("BR-AG-", "BR‑AG‑")
                      .replace("TC-AG-", "TC‑AG‑"))
        r.font.size = Pt(size)
        r.bold = bold


def table(headers, rows, widths, size=11, center=(), caption=None, label_col=False):
    if caption:
        p = para(caption, L, 12, after=3, indent=False, italic=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
    elif doc.paragraphs and doc.paragraphs[-1].text and not doc.paragraphs[-1].style.name.startswith("Heading"):
        doc.paragraphs[-1].paragraph_format.space_after = Pt(6)
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if headers:
        row = t.add_row()
        for i, text in enumerate(headers):
            fill(row.cells[i], text, size, bold=True, align=C)
            shade(row.cells[i], "D9D9D9")
        row_flag(row, "w:tblHeader")
    for data in rows:
        row = t.add_row()
        row_flag(row, "w:cantSplit")
        for i, text in enumerate(data):
            fill(row.cells[i], text, size, bold=label_col and i == 0, align=C if i in center else L)
            if label_col and i == 0:
                shade(row.cells[i], "F2F2F2")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(round(w / 2.54 * 1440)))
    doc.add_paragraph()
    return t


FIG = [0]


def figure(file, caption, width=16.5):
    FIG[0] += 1
    doc.add_picture(str(SHOTS / file), width=Cm(width))
    pic = doc.paragraphs[-1]
    pic.alignment = C
    pic.paragraph_format.keep_with_next = True
    pic.paragraph_format.space_before = Pt(6)
    para(f"Рисунок {FIG[0]} — {caption}", C, 12, after=10)
    return FIG[0]


bugs = BUGS["bugs"]
excel_bug = next(b for b in bugs if b["where"] == "excel")
jira_bugs = [b for b in bugs if b["where"] == "jira"]

# ---------------------------------------------------------------- title page
for line in ["Міністерство освіти і науки України", "Національний технічний університет України",
             "«Київський політехнічний інститут ім. І. Сікорського»",
             "Кафедра інженерії програмного забезпечення в енергетиці"]:
    para(line, C)
for _ in range(6):
    para(align=C)
para("Практична робота № 4", C, 18, bold=True)
para("з курсу: «Компоненти програмної інженерії. Частина 4. Якість та тестування "
     "програмного забезпечення»", C)
para(align=C)
para("Тема: «Розробка звіту про дефект (Bug Report)»", C, bold=True)
para(f"Модуль «Агент БПЛА» проєкту «{BACKLOG['project']['title']}»", C)
for _ in range(5):
    para(align=C)
for line in ["Виконав:", "студент 3-го курсу, групи ТВ-43", "Степаненко Назар Юрійович",
             "Перевірив:", "доц., к.т.н. Варава І.А."]:
    para(line, R)
for _ in range(5):
    para(align=C)
para("Київ 2026", C)
doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)

# ---------------------------------------------------------------- 1. object and method
h1("1. Об'єкт тестування та спосіб пошуку дефектів")
para("Для пошуку дефектів обрано модуль «Агент БПЛА» (епік UAVS-7) — той самий, для якого в "
     "практичній роботі № 3 розроблено тест-кейси TC-AG-01…TC-AG-09. Програмний код модуля станом на "
     f"{BUGS['date_reported']} ще не реалізовано: історії UAVS-18…UAVS-20 заплановано на спрінт 3. Тому "
     "об'єктом тестування є вимоги до модуля, записані в Jira, а дефекти знайдено статичним "
     "тестуванням — рецензуванням вимог без запуску програми.")
para("Такий підхід відповідає матеріалу курсу: експертиза вимог належить до методів верифікації, "
     "перевірку повноти, однозначності й несуперечливості вимог виконує тест-аналітик, а помилки у "
     "вимогах є одним із джерел дефектів, які потім найважче відстежити в коді. Виправлення вимоги "
     "до початку розробки коштує найменше.")
rows = []
for i in (5, 6, 7):
    s = BACKLOG["stories"][i]
    rows.append([f"UAVS-{STORY_KEY_OFFSET + i}", [s["summary"]] + [f"– {ac}" for ac in s["ac"]]])
table(["Історія", "Вимога та критерії приймання (текст у Jira)"], rows, [2.3, 14.2],
      caption="Таблиця 1.1 — Вимоги модуля «Агент БПЛА», що тестувалися")
para("Кожен критерій приймання перевірено за чотирма ознаками якості вимоги:")
bullets([
    ("повнота — ", "чи містить критерій усі значення, потрібні для реалізації;"),
    ("однозначність — ", "чи можна його прочитати лише одним способом;"),
    ("несуперечливість — ", "чи узгоджується він з іншими вимогами проєкту;"),
    ("тестованість — ", "чи можна з нього вивести очікуваний результат тест-кейса, зокрема для "
                        "граничних і негативних сценаріїв."),
])
para("Для тестованості використано тест-кейси ПР3: якщо очікуваний результат тест-кейса не "
     "випливає з вимоги, а мав бути прийнятий як припущення, це фіксувалося як дефект вимоги.")
figure("req-UAVS-19.png", "Історія UAVS-19 у Jira: критерій «При заряді < порогу» без значення порогу")

# ---------------------------------------------------------------- 2. jira setup
h1("2. Налаштування Jira для реєстрації дефектів")
para("На початку роботи простір UAVS не мав типу робочого елемента Bug, тож зареєструвати дефект як "
     "дефект було неможливо. Відповідно до тест-плану UAVS-MTP-1.0 (завдання T2) тип Bug додано та "
     "налаштовано так, щоб він містив усі атрибути баг-репорту з лекції 5:")
table(["Атрибут баг-репорту", "Поле в Jira"], [
    ["Summary (тема)", "Summary"],
    ["Description, Steps To Reproduce, Actual / Expected Result", "Description — розділи «Опис», "
                                                                 "«Передумови», «Кроки відтворення», "
                                                                 "«Фактичний результат», «Очікуваний "
                                                                 "результат», «Вплив»"],
    ["Оточення дефекту", "Environment (додано до типу Bug)"],
    ["Priority (пріоритет)", "Priority: High, Medium, Low (додано до типу Bug)"],
    ["Severity (серйозність)", "Severity — нове поле-список S1 Blocker … S7 Trivial"],
    ["Status (статус)", "Status: To Do → In Progress → In Review → Testing → Done"],
    ["Attachments (вкладення)", "Attachments — знімок екрана вимоги"],
    ["Зв'язок з вимогою", "Linked work items (relates to) та Parent — епік UAVS-7"],
], [6.0, 10.5], size=11, caption="Таблиця 2.1 — Атрибути баг-репорту та поля типу Bug")
figure("01-bug-worktype.png", "Налаштування типу Bug у просторі UAVS: поля Environment, Priority, Severity",
       width=14.0)

# ---------------------------------------------------------------- 3. found bugs
h1("3. Знайдені дефекти")
table(["№", "Дефект", "Вимога", "Серйозність", "Пріоритет", "Де оформлено"],
      [[str(n), b["summary"].split(": ", 1)[1], b["story"], b["severity"], b["priority"],
        "Excel-файл" if b["where"] == "excel" else f"Jira, {b['jira_key']}"]
       for n, b in enumerate(bugs, 1)],
      [0.8, 5.5, 2.2, 2.8, 2.4, 2.8], size=11, center=(0, 2, 3, 4, 5),
      caption="Таблиця 3.1 — Перелік знайдених дефектів")
para(f"Усього знайдено {len(bugs)} дефекти. Дефект із найбільшим впливом — відсутність значення "
     "порогу повернення на базу — оформлено як окремий баг-репорт в Excel; решту задокументовано в "
     "Jira. Серйозність визначено за шкалою S1–S7: S3 (Major) означає, що частину основної логіки "
     "модуля неможливо реалізувати однозначно, але існує обхідний шлях — прийняти значення як "
     "припущення; S4 (Minor) — розбіжність стосується лише граничного випадку.")

# ---------------------------------------------------------------- 4. excel bug
h1(f"4. Баг-репорт {excel_bug['id']} (Excel)")
para(f"Баг-репорт оформлено у файлі «Баг-репорт_{excel_bug['id']}.xlsx» (аркуш «Баг-репорт»). Поля "
     "Severity, Priority і Status заповнюються зі списку значень; на аркуші «Вкладення» розміщено "
     "знімок екрана вимоги, на аркуші «Довідник» — шкали серйозності та пріоритету.")
table(None, [
    ["Bug ID", excel_bug["id"]],
    ["Summary", excel_bug["summary"]],
    ["Severity / Priority", f"{excel_bug['severity']} / {excel_bug['priority']}"],
    ["Status", "New"],
    ["Оточення", BUGS["environment"]],
    ["Автор, дата", f"{BUGS['found_by']}; виявлено {BUGS['date_found']}, звіт {BUGS['date_reported']}"],
    ["Description", excel_bug["description"]],
    ["Передумови", excel_bug["preconditions"]],
    ["Steps To Reproduce", [f"{n}. {s}" for n, s in enumerate(excel_bug["steps"], 1)]],
    ["Actual Result", excel_bug["actual"]],
    ["Expected Result", excel_bug["expected"]],
    ["Вплив", excel_bug["impact"]],
    ["Attachments", "Знімок екрана історії UAVS-19 (рисунок 1)"],
    ["Пов'язане", excel_bug["related"]],
], [4.0, 12.5], size=11, label_col=True, caption=f"Таблиця 4.1 — Зміст баг-репорту {excel_bug['id']}")

# ---------------------------------------------------------------- 5. jira bugs
h1("5. Баг-репорти в Jira")
para("Решту дефектів зареєстровано в просторі UAVS як робочі елементи типу Bug. Кожен баг-репорт "
     "пов'язано з історією, у вимогах якої знайдено дефект, і з епіком UAVS-7; до кожного додано "
     "знімок екрана вимоги. Виконавця не призначено: за тест-планом новий дефект отримує виконавця "
     "від Product Owner.")
figure("bugs-list.png", "Список дефектів простору UAVS (тип Bug)")
for b in jira_bugs:
    h2(f"5.{jira_bugs.index(b) + 1}. {b['jira_key']}. {b['summary'].split(': ', 1)[1]}")
    para(f"{b['description']} Серйозність — {b['severity']}, пріоритет — {b['priority']}. {b['impact']}")
    figure(f"jira-{b['jira_key']}.png", f"Баг-репорт {b['jira_key']} у Jira: опис, кроки відтворення, "
                                         "результати, оточення та панель атрибутів")

# ---------------------------------------------------------------- conclusions
h1("Висновки")
para(f"Під час статичного тестування вимог модуля «Агент БПЛА» знайдено {len(bugs)} дефекти: три з "
     "них стосуються неповноти вимог (не задано поріг повернення, коефіцієнт витрати й швидкість "
     "зарядки, поведінку при недопустимій цілі), один — неоднозначності (межа радіуса сенсора). "
     "Усі дефекти виявлено до написання коду, тому їх виправлення обмежується уточненням "
     "критеріїв приймання в Jira.")
para("Для реєстрації дефектів у Jira додано тип Bug із полями Severity, Priority та Environment, "
     "що відповідають атрибутам баг-репорту з лекції. Один дефект оформлено як баг-репорт в Excel, "
     "три інші — у Jira з кроками відтворення, фактичним і очікуваним результатом та вкладеннями.")

out = HERE / "Протокол_ПР4_Степаненко.docx"
doc.save(out)
print(f"{out.name}: {len(bugs)} bugs, {FIG[0]} figures, {len(doc.tables)} tables")
