"""Builds Баг-репорт_BR-AG-01.xlsx — the Excel bug report for the one bug not filed in Jira.

Content comes from bugs.json (the same source the Jira issues and the protocol use).
Attributes follow lecture 5: Summary, Description, Steps To Reproduce, Actual/Expected,
Attachments, Priority, Severity, Status, Environment.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
DATA = json.loads((HERE / "bugs.json").read_text(encoding="utf-8"))
BUG = next(b for b in DATA["bugs"] if b["where"] == "excel")

FONT = "Times New Roman"
thin = Side(style="thin", color="808080")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LABEL_FILL = PatternFill("solid", fgColor="D9D9D9")
SECTION_FILL = PatternFill("solid", fgColor="BDD7EE")
WRAP = Alignment(wrap_text=True, vertical="top")

SEVERITIES = ["S1 Blocker", "S2 Critical", "S3 Major", "S4 Minor", "S5 Tweak", "S6 Text", "S7 Trivial"]
PRIORITIES = ["High", "Medium", "Low"]
STATUSES = ["New", "Assigned", "Open", "Fixed", "Verified", "Closed", "Reopened", "Rejected", "Deferred"]

wb = Workbook()
ws = wb.active
ws.title = "Баг-репорт"
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 95
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:B1")
ws["A1"] = "Звіт про дефект (Bug Report)"
ws["A1"].font = Font(name=FONT, size=16, bold=True)
ws["A1"].alignment = Alignment(horizontal="center")
ws.row_dimensions[1].height = 26
ws.merge_cells("A2:B2")
ws["A2"] = "Проєкт UAV Swarm Mesa (Jira: UAVS) · Курс «Якість та тестування ПЗ» · ПР № 4"
ws["A2"].font = Font(name=FONT, size=11, italic=True)
ws["A2"].alignment = Alignment(horizontal="center")

steps = "\n".join(f"{n}. {s}" for n, s in enumerate(BUG["steps"], 1))
sections = [
    ("Ідентифікація", [
        ("Bug ID", BUG["id"]),
        ("Summary (тема)", BUG["summary"]),
        ("Проєкт", "Мультиагентна програмна система моделювання та координації групи БПЛА на основі Mesa"),
        ("Модуль", "Агент БПЛА (епік UAVS-7)"),
        ("Об'єкт тестування", "Вимога UAVS-19, критерій «При заряді < порогу стан RETURNING»"),
    ]),
    ("Класифікація", [
        ("Severity (серйозність)", BUG["severity"]),
        ("Priority (пріоритет)", BUG["priority"]),
        ("Status (статус)", "New"),
        ("Тип дефекту", "Дефект вимог: неповнота (відсутнє числове значення) і неоднозначність"),
    ]),
    ("Оточення та люди", [
        ("Оточення", DATA["environment"]),
        ("Автор (Reporter)", DATA["found_by"]),
        ("Виконавець (Assignee)", "Не призначено — призначає Product Owner"),
        ("Дата виявлення", f"{DATA['date_found']} — {DATA['how_found']}"),
        ("Дата звіту", DATA["date_reported"]),
    ]),
    ("Опис дефекту", [
        ("Description (опис)", BUG["description"]),
        ("Передумови", BUG["preconditions"]),
        ("Steps To Reproduce", steps),
        ("Actual Result (фактичний)", BUG["actual"]),
        ("Expected Result (очікуваний)", BUG["expected"]),
        ("Вплив", BUG["impact"]),
    ]),
    ("Додатково", [
        ("Attachments (вкладення)", "Знімок екрана історії UAVS-19 — аркуш «Вкладення»"),
        ("Пов'язані вимоги й тести", BUG["related"]),
        ("Примітка", "Інші дефекти, знайдені під час цього тестування, задокументовано в Jira: "
                     + ", ".join(f"{b['jira_key']} ({b['severity']})" for b in DATA["bugs"] if b.get("jira_key"))),
    ]),
]

row = 4
cells = {}
for title, fields in sections:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name=FONT, size=12, bold=True)
    c.fill = SECTION_FILL
    for col in (1, 2):
        ws.cell(row=row, column=col).border = BOX
    row += 1
    for label, value in fields:
        a = ws.cell(row=row, column=1, value=label)
        b = ws.cell(row=row, column=2, value=value)
        a.font = Font(name=FONT, size=11, bold=True)
        a.fill = LABEL_FILL
        b.font = Font(name=FONT, size=11)
        for c in (a, b):
            c.border = BOX
            c.alignment = WRAP
        lines = sum(max(1, -(-len(part) // 85)) for part in str(value).split("\n"))   # ~85 chars per line at width 95
        ws.row_dimensions[row].height = max(18, 15 * lines + 4)
        cells[label] = b.coordinate
        row += 1
    row += 1

# keep the classification cells to the scales from lecture 5
for label, values in (("Severity (серйозність)", SEVERITIES), ("Priority (пріоритет)", PRIORITIES),
                      ("Status (статус)", STATUSES)):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    dv.error, dv.errorTitle = "Оберіть значення зі списку", "Недопустиме значення"
    ws.add_data_validation(dv)
    dv.add(cells[label])
ws.print_area = f"A1:B{row}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# attachments sheet: the requirement as it reads in Jira
att = wb.create_sheet("Вкладення")
att["A1"] = "Вкладення 1 — історія UAVS-19 у Jira станом на 21.10.2026: критерій посилається на «поріг», значення якого не задано"
att["A1"].font = Font(name=FONT, size=11, bold=True)
img = XLImage(str(HERE / "screenshots" / BUG["attachments"][0]))
img.width, img.height = int(img.width * 0.7), int(img.height * 0.7)
att.add_image(img, "A3")

# reference scales
ref = wb.create_sheet("Довідник")
ref.column_dimensions["A"].width = 16
ref.column_dimensions["B"].width = 90
rows = [("Severity", "Вплив дефекту на працездатність (лекція 5, тест-план UAVS-MTP-1.0, додаток Б)"),
        ("S1 Blocker", "Система не працює, подальше тестування неможливе"),
        ("S2 Critical", "Не працює ключова логіка, обхідного шляху немає"),
        ("S3 Major", "Частина основної логіки працює неправильно, є обхідний шлях"),
        ("S4 Minor", "Не порушує бізнес-логіку; помітна помилка"),
        ("S5 Tweak", "Незручність інтерфейсу"),
        ("S6 Text", "Текстова помилка або одрук"),
        ("S7 Trivial", "Не впливає на якість продукту"),
        ("", ""),
        ("Priority", "Терміновість виправлення; встановлює Product Owner"),
        ("High", "Виправити якнайшвидше"),
        ("Medium", "Виправити обов'язково, можна в наступному спрінті"),
        ("Low", "Виправлення можна відкласти")]
for r, (a, b) in enumerate(rows, 1):
    ref.cell(row=r, column=1, value=a).font = Font(name=FONT, size=11, bold=a in ("Severity", "Priority"))
    ref.cell(row=r, column=2, value=b).font = Font(name=FONT, size=11, bold=a in ("Severity", "Priority"))

out = HERE / "Баг-репорт_BR-AG-01.xlsx"
wb.save(out)
print(f"{out.name}: {row - 4} rows on the report sheet, attachment {BUG['attachments'][0]}")
