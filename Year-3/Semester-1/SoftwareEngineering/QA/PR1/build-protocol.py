"""Builds Протокол_ПР1.docx from backlog.json and screenshots/*.png."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
DATA = json.loads((HERE / "backlog.json").read_text(encoding="utf-8"))
SHOTS = HERE / "screenshots"

# (file, caption) in the order the assignment lists its steps
SCREENSHOTS = [
    ("01-project.png", "Сторінка Summary проєкту UAV Swarm Mesa у Jira"),
    ("02-people.png", "Учасники проєкту (Space settings → Access)"),
    ("03-timeline.png", "План виконання проєкту на семестр: епіки та 7 спрінтів на Timeline"),
    ("04-backlog.png", "Беклог: user stories з епіками, оцінками та виконавцями, розподілені по спрінтах"),
    ("05-story.png", "User story UAVS-16 з критеріями приймання, пріоритетом і story points"),
    ("07-board-settings.png", "Налаштування колонок і статусів дошки (додано етап Testing)"),
    ("08-board.png", "Scrum-дошка запущеного спрінту 1"),
]
EPIC_KEY_OFFSET = 5  # epics were created right after 4 deleted template issues
SITE = "https://tv43-uav-swarm.atlassian.net"

doc = Document()
sec = doc.sections[0]
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)

base = doc.styles["Normal"]
base.font.name = "Times New Roman"
base.font.size = Pt(14)
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
base.paragraph_format.space_after = Pt(0)
base.paragraph_format.line_spacing = 1.15


def para(text="", align=None, size=None, bold=False, space_after=0, indent=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return p


def heading(text, new_page=False):
    p = para(text, WD_ALIGN_PARAGRAPH.CENTER, 16, bold=True, space_after=8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = new_page


C, R, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY

# --- Титульна сторінка ---
for line in ["Міністерство освіти і науки України",
             "Національний технічний університет України",
             "«Київський політехнічний інститут ім. І. Сікорського»",
             "Кафедра інженерії програмного забезпечення в енергетиці"]:
    para(line, C)
for _ in range(6):
    para()
para("Практична робота № 1", C, 18, bold=True)
para("з курсу: «Компоненти програмної інженерії. Частина 4. Якість та тестування програмного забезпечення»", C)
para()
para("Тема: «Система управління розробкою як інструмент управління якістю»", C, bold=True)
for _ in range(6):
    para()
para("Виконали:", R)
para("студенти 3-го курсу, групи ТВ-43", R)
for m in DATA["team"]:
    para(m["name"], R)
para("Перевірив:", R)
para("Варава І.", R)
for _ in range(5):
    para()
para("Київ 2026", C)
doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)

# --- 1. Команда ---
heading("1. Склад команди та ролі")
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(table.rows[0].cells, ["№", "ПІБ", "Ролі в проєкті"]):
    cell.text = text
    cell.paragraphs[0].runs[0].bold = True
for i, m in enumerate(DATA["team"], 1):
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = str(i), m["name"], ", ".join(m["roles"])
para(space_after=6)
para("Відповідно до вимог курсу кожен учасник виконує як роль розробника, так і роль "
     "тестувальника: задачі на реалізацію та перевірку розподілено перехресно, щоб "
     "код одного учасника тестував і рецензував інший.", J, indent=True)

# --- 2. Опис проєкту ---
heading("2. Короткий опис проєкту")
para(f"Тема: «{DATA['project']['title']}».", J, indent=True)
para(DATA["project"]["description"], J, indent=True)
para(f"Jira: {SITE}, ключ проєкту {DATA['project']['key']}. Шаблон: Scrum (team-managed). "
     f"Технології: Python 3.12, Mesa 3, NumPy, Solara, pytest, GitHub Actions.", J, indent=True)

# --- 3. План ---
heading("3. План виконання проєкту на семестр")
para("Робота ведеться двотижневими спрінтами. Епіки мають дати початку й завершення, "
     "що формує план на вкладці Timeline.", J, indent=True, space_after=6)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
for cell, text in zip(t.rows[0].cells, ["Спрінт", "Період", "Мета спрінту"]):
    cell.text = text
    cell.paragraphs[0].runs[0].bold = True
for s in DATA["sprints"]:
    r = t.add_row().cells
    r[0].text = s["name"].replace("UAVS ", "")
    r[1].text = f"{s['start'][8:]}.{s['start'][5:7]} – {s['end'][8:]}.{s['end'][5:7]}"
    r[2].text = s["goal"]

# --- 4-5. Вимоги та беклог ---
heading("4. Вимоги до ПЗ та беклог")
names = {m["id"]: m["name"].split()[0] + " " + m["name"].split()[1][0] + "." for m in DATA["team"]}
for i, e in enumerate(DATA["epics"]):
    stories = [s for s in DATA["stories"] if s["epic"] == e["key"]]
    para(f"Епік {DATA['project']['key']}-{i + EPIC_KEY_OFFSET}. {e['summary']} "
         f"({sum(s['sp'] for s in stories)} SP)", bold=True, space_after=2)
    for s in stories:
        p = para(f"• {s['summary']} — {s['sp']} SP, спрінт {s['sprint']}, {names[s['assignee']]}", J)
        p.paragraph_format.left_indent = Cm(0.75)
        p.runs[0].font.size = Pt(12)
    para(space_after=4)
total = sum(s["sp"] for s in DATA["stories"])
para(f"Разом: {len(DATA['epics'])} епіків, {len(DATA['stories'])} user stories ({total} story points), "
     f"{len(DATA['tasks'])} підзадач спрінту 1.", J, indent=True)

# --- 6. Скріншоти ---
heading("5. Робота в Jira", new_page=True)
n = 0
for file, caption in SCREENSHOTS:
    path = SHOTS / file
    if not path.exists():
        continue
    n += 1
    doc.add_picture(str(path), width=Cm(16.5))
    doc.paragraphs[-1].alignment = C
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(f"Рисунок {n} — {caption}", C, 12, space_after=10)

heading("Висновки")
para("У ході роботи створено проєкт у системі управління розробкою Jira, сформовано "
     "команду з розподілом ролей, зібрано вимоги у вигляді user stories з критеріями "
     "приймання, декомпоновано їх на епіки та задачі з оцінками і виконавцями, "
     "сплановано сім двотижневих спрінтів на семестр, налаштовано дошку з окремим "
     "етапом тестування та запущено перший спрінт. Критерії приймання та етап "
     "Testing на дошці роблять перевірку якості обов'язковою частиною процесу, "
     "а не окремою фазою наприкінці семестру.", J, indent=True)

out = HERE / "Протокол_ПР1.docx"
doc.save(out)
print(f"{out.name}: {n} screenshots embedded")
